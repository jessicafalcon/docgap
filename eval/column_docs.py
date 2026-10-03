"""Write the every-column docs: each mart column's text from the dictionary, verbatim in French.

A column named for a dictionary variable gets the variable's label (`Libellé`), then
a newline and its comment (`Commentaires`) when it has one. A `<VAR>_LIB` label
column, which the variable sheet doesn't list, gets the label cell of the line its
code list opens with: the text describing the labels the dimension holds, the one
its seed's labels are graded against. Each cell is stripped of leading and trailing
whitespace and otherwise copied as it is, so no wording is chosen (ADR 0025). The
text is the pilot's full-docs configuration and the ceiling arm's, and its locked
half is the baseline's (ADR 0021): it is frozen at `preregistered`.

    uv run python -m eval.column_docs
"""

from __future__ import annotations

import json
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path

import openpyxl
import yaml

__all__ = [
    "CODE_LIST_SHEET",
    "DICTIONARY",
    "DOCS",
    "MARTS_YAML",
    "VARIABLE_SHEET",
    "build",
    "column_docs",
    "label_texts",
    "mart_columns",
    "render",
    "variable_texts",
]

ROOT = Path(__file__).resolve().parents[1]
# The seeds read the same file and sheets; a test checks the names agree.
DICTIONARY = ROOT / "eval" / "reference" / "2024_descriptif-variables_open-damir-base-complete.xlsx"
VARIABLE_SHEET = "OPEN DAMIR"
CODE_LIST_SHEET = "MOD OPEN DAMIR"
MARTS_YAML = ROOT / "warehouse" / "dbt" / "models" / "marts" / "_marts__models.yml"
DOCS = ROOT / "eval" / "column_docs.json"

_NAME, _LABEL, _COMMENT = "Nom variable", "Libellé", "Commentaires"


def _text(cell: object) -> str | None:
    return (cell.strip() or None) if isinstance(cell, str) else None


def variable_texts(rows: Iterable[Sequence[object]]) -> dict[str, str]:
    """Map each variable of the variable sheet to its label, then its comment if any.

    >>> variable_texts([
    ...     ("Nom variable", "Libellé", "Catégorie", "Commentaires"),
    ...     ("SOI_ANN", " Année de soins ", "PRESTATION", None),
    ...     ("ASU_NAT", "Nature d'Assurance", "PRESTATION", "Maladie,\\nmaternité "),
    ... ])
    {'SOI_ANN': 'Année de soins', 'ASU_NAT': "Nature d'Assurance\\nMaladie,\\nmaternité"}

    Raises:
        ValueError: no header line, a variable with no label, or a variable twice.
    """
    texts: dict[str, str] = {}
    columns: dict[str, int] | None = None
    for row in rows:
        if columns is None:
            if _NAME in row:
                columns = {str(cell): i for i, cell in enumerate(row) if isinstance(cell, str)}
            continue
        name = _text(row[columns[_NAME]]) if len(row) > columns[_NAME] else None
        if name is None:
            continue
        label, comment = (
            _text(row[columns[key]]) if len(row) > columns[key] else None
            for key in (_LABEL, _COMMENT)
        )
        if label is None:
            raise ValueError(f"{VARIABLE_SHEET}: {name} has no label")
        if name in texts:
            raise ValueError(f"{VARIABLE_SHEET}: {name} listed twice")
        texts[name] = label if comment is None else f"{label}\n{comment}"
    if columns is None:
        raise ValueError(f"{VARIABLE_SHEET}: no line names its columns {_NAME!r}")
    return texts


def label_texts(rows: Iterable[Sequence[object]]) -> dict[str, str]:
    """Map `<VAR>_LIB` to the label cell of the line opening `VAR`'s code list.

    A list opens with a line whose first cell is the variable's name; its codes are
    integers. The sheet's title line opens nothing, since its second cell is empty.

    >>> label_texts([
    ...     ("MODALITES VARIABLES OPEN DAMIR", None),
    ...     ("PRS_NAT", "Libellé Nature de Prestation ", "Libellé en B2"),
    ...     (0, "SANS OBJET"),
    ... ])
    {'PRS_NAT_LIB': 'Libellé Nature de Prestation'}
    """
    texts: dict[str, str] = {}
    for row in rows:
        name = _text(row[0]) if row else None
        label = _text(row[1]) if len(row) > 1 else None
        if name is not None and label is not None:
            texts[f"{name}_LIB"] = label
    return texts


def mart_columns(marts_yaml: Path = MARTS_YAML) -> list[str]:
    """Every column name the marts YAML declares, once each, sorted."""
    models = yaml.safe_load(marts_yaml.read_text())["models"]
    return sorted({column["name"] for model in models for column in model.get("columns", [])})


def column_docs(
    columns: Iterable[str],
    variables: Mapping[str, str],
    labels: Mapping[str, str],
) -> dict[str, str]:
    """Give each column its dictionary text: a variable's, else a code list's label line.

    Raises:
        ValueError: a column has no text in the dictionary.
    """
    docs = {column: variables.get(column) or labels.get(column) for column in columns}
    if missing := sorted(column for column, text in docs.items() if text is None):
        raise ValueError(f"no dictionary text for the mart columns {missing}")
    return {column: text for column, text in sorted(docs.items()) if text is not None}


def render(docs: Mapping[str, str]) -> bytes:
    """Return the docs file: JSON sorted by column, non-ASCII kept as it is."""
    return (json.dumps(dict(sorted(docs.items())), ensure_ascii=False, indent=2) + "\n").encode()


def build(dictionary: Path = DICTIONARY, marts_yaml: Path = MARTS_YAML) -> bytes:
    """Return the docs file's bytes, read from the dictionary and the marts YAML."""
    workbook = openpyxl.load_workbook(dictionary, read_only=True, data_only=True)
    try:
        variables = variable_texts(workbook[VARIABLE_SHEET].iter_rows(values_only=True))
        labels = label_texts(workbook[CODE_LIST_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    return render(column_docs(mart_columns(marts_yaml), variables, labels))


def main() -> None:
    """Write the docs file beside this module."""
    content = build()
    DOCS.write_bytes(content)
    print(f"{DOCS.relative_to(ROOT).as_posix()}: {len(json.loads(content))} columns")


if __name__ == "__main__":
    main()
